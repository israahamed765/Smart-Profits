"""Smart Profits — GSMA MENA Ignite comprehensive pitch deck (16:9 PPTX)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x11, 0x20)
CYAN = RGBColor(0x4F, 0xD1, 0xC5)
CYAN_D = RGBColor(0x0F, 0x9E, 0x94)
GOLD = RGBColor(0xE8, 0xC5, 0x6B)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
CARD = RGBColor(0x16, 0x21, 0x36)
GREEN = RGBColor(0x34, 0xD3, 0x99)
RED = RGBColor(0xF8, 0x71, 0x71)
AMBER = RGBColor(0xFB, 0xBF, 0x24)

W = Inches(13.333)
H = Inches(7.5)
ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
LOGO = ROOT.parent.parent / "public" / "brand" / "mark.png"

TOTAL = 18
page = 0

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


def set_run(run, size=18, bold=False, color=WHITE, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(sh, color)
    return sh


def add_round(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(sh, color)
    sh.adjustments[0] = 0.08
    return sh


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def add_paras(slide, l, t, w, h, lines, size=14, color=WHITE, bullet=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        if bullet:
            p.level = 0
        run = p.add_run()
        run.text = ("• " if bullet else "") + line
        set_run(run, size, False, color)
    return box


def footer(slide):
    global page
    page += 1
    add_text(slide, Inches(0.5), Inches(7.12), Inches(9), Inches(0.28), "Smart Profits  ·  GSMA MENA Ignite  ·  Theme 4  ·  Nokia Network-as-Code", 10, False, MUTED)
    add_text(slide, Inches(11.4), Inches(7.12), Inches(1.4), Inches(0.28), f"{page:02d} / {TOTAL:02d}", 10, False, MUTED, PP_ALIGN.RIGHT)


def chrome(slide, kicker, title, subtitle=None):
    add_rect(slide, 0, 0, W, H, NAVY)
    add_rect(slide, 0, 0, W, Inches(0.07), CYAN)
    add_rect(slide, 0, Inches(0.07), W, Inches(0.015), GOLD)
    add_text(slide, Inches(0.55), Inches(0.28), Inches(12), Inches(0.32), kicker.upper(), 11, True, CYAN)
    add_text(slide, Inches(0.55), Inches(0.58), Inches(12.2), Inches(0.65), title, 26, True, WHITE)
    if subtitle:
        add_text(slide, Inches(0.55), Inches(1.18), Inches(12.2), Inches(0.45), subtitle, 14, False, MUTED)


def card(slide, l, t, w, h, heading, body, heading_color=GOLD, body_size=13):
    add_round(slide, l, t, w, h, CARD)
    add_text(slide, l + Inches(0.22), t + Inches(0.14), w - Inches(0.4), Inches(0.34), heading, 15, True, heading_color)
    add_text(slide, l + Inches(0.22), t + Inches(0.48), w - Inches(0.4), h - Inches(0.62), body, body_size, False, WHITE)


def badge_row(slide, y, tags, color=RGBColor(0x14, 0x3A, 0x3A)):
    w_each = Inches(12.3 / max(len(tags), 1))
    for i, tag in enumerate(tags):
        x = Inches(0.5) + i * w_each
        add_round(slide, x, y, w_each - Inches(0.12), Inches(0.42), color)
        add_text(slide, x, y + Inches(0.04), w_each - Inches(0.12), Inches(0.36), tag, 11, True, CYAN, PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# 01 Title
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, 0, 0, W, Inches(0.08), CYAN)
add_rect(s, 0, Inches(0.08), W, Inches(0.02), GOLD)
if LOGO.exists():
    s.shapes.add_picture(str(LOGO), Inches(0.55), Inches(1.15), Inches(1.1), Inches(1.1))
add_text(s, Inches(0.55), Inches(2.35), Inches(12), Inches(0.32), "GSMA MENA IGNITE OPEN GATEWAY HACKATHON  ·  2026", 12, True, CYAN)
add_text(s, Inches(0.55), Inches(2.72), Inches(12.2), Inches(0.85), "Smart Profits", 46, True, WHITE)
add_text(s, Inches(0.55), Inches(3.55), Inches(12), Inches(0.55), "AI-Driven Financial Analytics & Telco-Secured Merchant Platform", 21, False, GOLD)
add_text(
    s, Inches(0.55), Inches(4.2), Inches(11.8), Inches(0.9),
    "Theme 4: Secure FinTech, Payments & Anti-Fraud Innovation\n"
    "The merchant’s CFO in the chat. Bank-grade SIM-swap defence on the login.\n"
    "CAMARA APIs via Nokia Network-as-Code. Decisions by an AI Agent — not by SMS.",
    15, False, MUTED,
)
badge_row(s, Inches(5.35), ["SIM Swap", "Number Verification", "Location Verification", "Device Swap", "Smart Guard AI Agent"])
add_text(s, Inches(0.55), Inches(6.05), Inches(12), Inches(0.55),
         "Israa Nael Hamad  ·  University of Palestine, Gaza  ·  Bitsandbytesdude Software Agency  ·  MENA",
         13, False, GOLD)
footer(s)

# ---------------------------------------------------------------------------
# 02 Executive summary (judges scan this first)
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "Executive summary", "One slide for the jury", "What we built, why it wins Theme 4, and how both mandatory requirements are met.")
cols = [
    ("The problem", "MENA merchants run the business from Excel. Phantom profit hides real cash flow. A SIM swap defeats SMS OTP — then the attacker owns the P&L files, prices, and exports."),
    ("Our solution", "Smart Profits = Financial AI + Smart Guard AI Agent. Analytics on messy Arabic/English ledgers, protected by GSMA CAMARA signals from Nokia Network-as-Code."),
    ("Why we win", "Real product (Next.js + PostgreSQL). Real agent policy (Allow / Step-up / Freeze). Live Nokia RapidAPI logs in prototype. Built for MENA (RTL Arabic + English)."),
]
for i, (h, b) in enumerate(cols):
    card(s, Inches(0.5) + i * Inches(4.2), Inches(1.75), Inches(3.95), Inches(2.55), h, b)
add_round(s, Inches(0.5), Inches(4.55), Inches(12.3), Inches(1.95), CARD)
add_text(s, Inches(0.75), Inches(4.7), Inches(12), Inches(0.35), "Mandatory requirements — explicit declaration", 15, True, GOLD)
add_paras(s, Inches(0.75), Inches(5.05), Inches(11.8), Inches(1.3), [
    "Requirement A — GSMA CAMARA APIs via Nokia Network-as-Code: SIM Swap, Number Verification, Location Verification, Device Swap on every sensitive merchant action.",
    "Requirement B — AI Agent orchestration: Smart Guard calls network tools, fuses risk with financial context, and executes Allow / Step-up / Freeze — not a chatbot caption.",
], 13, WHITE, bullet=True)
footer(s)

# ---------------------------------------------------------------------------
# 03 Problem
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "1. Problem statement", "The P&L lives in Excel. So does the fraud.")
card(s, Inches(0.5), Inches(1.45), Inches(4.0), Inches(3.2), "Financial blindness",
     "Independent MENA merchants use messy Arabic/English Excel, CSV, PDF and photos. Rent, salaries and utilities are missing. Phantom profit looks healthy. There is no CFO.")
card(s, Inches(4.65), Inches(1.45), Inches(4.0), Inches(3.2), "Profit leakage",
     "Weak SKUs, unexamined margins and unplanned OpEx hide real net profit. A ‘good month’ in the sales sheet can still be cash-negative.")
card(s, Inches(8.8), Inches(1.45), Inches(4.0), Inches(3.2), "Account takeover",
     "SIM swap beats SMS OTP. Attacker opens advisor, downloads P&L, changes prices, exfiltrates sales files. Upload from another country = hijacked session.")
add_round(s, Inches(0.5), Inches(4.85), Inches(12.3), Inches(1.75), CARD)
add_text(s, Inches(0.75), Inches(5.0), Inches(11.8), Inches(1.45),
         "Theme 4 gap: Banks already buy Open Gateway APIs. The SaaS that holds the merchant’s books usually does not.\n"
         "Smart Profits closes that gap — financial intelligence and telco-grade identity on the same login.",
         15, False, WHITE)
footer(s)

# ---------------------------------------------------------------------------
# 04 MENA market & opportunity
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "2. Market & opportunity", "Why MENA, why now")
stats = [
    ("Excel-first commerce", "WhatsApp + spreadsheet workflows dominate SME retail. Arabic/English mixed headers are the norm, not the exception."),
    ("SIM-swap fraud", "OTP hijacking is a documented FinTech risk. Network verification beats SMS that can be intercepted after a swap."),
    ("Open Gateway moment", "GSMA CAMARA standardises operator APIs. Nokia Network-as-Code gives developers one path — Smart Profits is a reference FinTech consumer."),
    ("Operator B2B2X", "Every secured login, upload and export is a billable CAMARA event — SaaS revenue for us, API usage for operators."),
]
for i, (h, b) in enumerate(stats):
    col, row = i % 2, i // 2
    card(s, Inches(0.5) + col * Inches(6.4), Inches(1.45) + row * Inches(2.45), Inches(6.05), Inches(2.2), h, b)
footer(s)

# ---------------------------------------------------------------------------
# 05 Solution overview
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "3. Solution overview", "One platform. Two brains. Network-secured decisions.")
card(s, Inches(0.5), Inches(1.45), Inches(6.05), Inches(2.0), "Financial AI engine",
     "Cleans messy Excel/CSV/PDF. Diagnoses store health, profit leaks, inventory, 30-day plan. File-grounded Q&A in Arabic or English. OpEx-aware net profit — not phantom revenue.")
card(s, Inches(6.75), Inches(1.45), Inches(6.05), Inches(2.0), "Smart Guard AI Agent",
     "Orchestrates Nokia NaC CAMARA calls on login, reset, upload, export, price change. Outputs Allow · Step-up · Freeze. Explains decisions to the merchant in AR/EN.")
add_round(s, Inches(0.5), Inches(3.65), Inches(12.3), Inches(2.85), CARD)
add_text(s, Inches(0.75), Inches(3.8), Inches(12), Inches(0.35), "Sensitive actions protected by Smart Guard", 15, True, GOLD)
actions = "Register  ·  Login  ·  Password reset  ·  Excel/CSV/PDF upload  ·  P&L export  ·  Price / settings change"
add_text(s, Inches(0.75), Inches(4.2), Inches(11.8), Inches(0.45), actions, 14, True, CYAN)
add_text(s, Inches(0.75), Inches(4.75), Inches(11.8), Inches(1.55),
         "Positioning: Not a generic Excel AI. A FinTech control plane for SMEs — using 4G/5G CAMARA signals the way banks do, on the merchant’s own books.\n"
         "No sensitive action completes unless Smart Guard runs the network + policy loop.",
         14, False, WHITE)
footer(s)

# ---------------------------------------------------------------------------
# 06 Mandatory requirements mapping (critical for judges)
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "4. Mandatory hackathon requirements", "Both compulsory building blocks — implemented, not slideware")
headers = ["Requirement", "Implementation in Smart Profits", "Evidence"]
rows = [
    ("A — CAMARA via Nokia NaC",
     "SIM Swap, Number Verification, Location Verification, Device Swap via RapidAPI passthrough paths",
     "Live logs: [nac-live] provider=Nokia transport=RapidAPI status=200"),
    ("B — AI Agent orchestration",
     "Smart Guard Agent: tool-calling loop → policy.ts → Allow / Step-up / Freeze",
     "Server-side enforce.ts blocks API until decision=allow"),
    ("Theme 4 — FinTech anti-fraud",
     "Freeze on swap; step-up on NV/location gaps; block upload outside store geofence",
     "Demo MSISDNs + working UI overlay + admin guard status"),
]
table = s.shapes.add_table(4, 3, Inches(0.5), Inches(1.55), Inches(12.3), Inches(2.85)).table
col_w = [Inches(2.8), Inches(5.2), Inches(4.3)]
for j, h in enumerate(headers):
    table.columns[j].width = col_w[j]
    cell = table.cell(0, j)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0x0B, 0x11, 0x20)
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            set_run(r, 11, True, CYAN)
for i, row in enumerate(rows, start=1):
    for j, val in enumerate(row):
        cell = table.cell(i, j)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD if i % 2 else RGBColor(0x14, 0x1C, 0x30)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_run(r, 11, j == 0, WHITE if j else CYAN)
add_round(s, Inches(0.5), Inches(4.65), Inches(12.3), Inches(1.95), CARD)
add_text(s, Inches(0.75), Inches(4.8), Inches(12), Inches(0.35), "Declaration", 14, True, GOLD)
add_text(s, Inches(0.75), Inches(5.15), Inches(11.8), Inches(1.25),
         "This solution uses GSMA CAMARA network APIs through Nokia Network-as-Code AND uses an AI Agent to orchestrate those APIs and take security decisions. Both mandatory requirements are satisfied in the live decision path — not as optional plugins.",
         13, False, WHITE)
footer(s)

# ---------------------------------------------------------------------------
# 07 CAMARA APIs detail
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "5. Nokia CAMARA APIs", "What we call, when, and what the agent does")
headers = ["CAMARA API", "Trigger", "Agent decision"]
rows = [
    ("SIM Swap", "Login, register, reset, sensitive settings", "Recent swap → FREEZE session. SMS OTP is untrusted."),
    ("Number Verification", "Login, step-up identity", "Silent 4G/5G check — confirm live network number, not spoofed SMS."),
    ("Location Verification", "Financial file upload, P&L export", "Must match store geofence — else STEP-UP or block upload."),
    ("Device Swap", "Same sensitive actions as SIM", "Recent device change → FREEZE (complements SIM signal)."),
]
table = s.shapes.add_table(5, 3, Inches(0.5), Inches(1.5), Inches(12.3), Inches(3.1)).table
for j, h in enumerate(headers):
    table.columns[j].width = col_w[j]
    cell = table.cell(0, j)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0x0B, 0x11, 0x20)
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            set_run(r, 11, True, CYAN)
for i, row in enumerate(rows, start=1):
    for j, val in enumerate(row):
        cell = table.cell(i, j)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD if i % 2 else RGBColor(0x14, 0x1C, 0x30)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_run(r, 11, j == 0, WHITE if j else CYAN)
add_text(s, Inches(0.55), Inches(4.85), Inches(12.2), Inches(1.5),
         "Prototype: Nokia NaC developer simulators + live RapidAPI (NAC_API_KEY).\n"
         "Official test MSISDNs: +99999991001 clean · +99999991000 swap/freeze · +99999991002 location step-up · +99999991003 location unknown.\n"
         "Policy engine (policy.ts) is unchanged — Nokia results map through nac-live-mapper; no hardcoded fake success on provider failure.",
         12, False, MUTED)
footer(s)

# ---------------------------------------------------------------------------
# 08 AI Agent design
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "6. Smart Guard AI Agent", "Sense → Query → Decide → Act")
steps = [
    ("1 · Sense", "Merchant action: login, register, upload, export, price change."),
    ("2 · Query", "Parallel CAMARA calls via Nokia NaC: SIM Swap, Device Swap, NV, Location."),
    ("3 · Decide", "Fuse network risk + file sensitivity + account age. Output: Allow · Step-up · Freeze."),
    ("4 · Act", "Enforce on server (403/503). UI overlay for step-up network code. Explain in AR/EN."),
]
for i, (title, body) in enumerate(steps):
    x = Inches(0.5) + i * Inches(3.2)
    card(s, x, Inches(1.45), Inches(3.0), Inches(2.15), title, body, CYAN, 12)
decisions = [
    ("ALLOW", "Signals consistent. Continue analytics / upload.", GREEN),
    ("STEP-UP", "NV unavailable or soft location mismatch. Network code — not SMS.", AMBER),
    ("FREEZE", "SIM/device swap or hard location mismatch. Block sensitive actions.", RED),
]
for i, (title, body, col) in enumerate(decisions):
    x = Inches(0.5) + i * Inches(4.2)
    add_round(s, x, Inches(3.85), Inches(3.95), Inches(1.55), CARD)
    add_text(s, x + Inches(0.2), Inches(3.98), Inches(3.5), Inches(0.35), title, 18, True, col)
    add_text(s, x + Inches(0.2), Inches(4.38), Inches(3.5), Inches(0.85), body, 12, False, WHITE)
add_text(s, Inches(0.55), Inches(5.65), Inches(12.2), Inches(1.2),
         "Example freeze: upload monthly_sales.xlsx → SIM Swap recent OR location ≠ store → FREEZE → file not stored → recover via Number Verification on trusted line.",
         13, False, MUTED)
footer(s)

# ---------------------------------------------------------------------------
# 09 Architecture
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "7. Technical architecture", "Production-ready stack — not a hackathon script")
boxes = [
    (0.5, 1.42, 12.3, 1.05, "Merchant UI  ·  Next.js / TypeScript  ·  Arabic RTL + English  ·  Dashboard · Advisor · Upload · Admin"),
    (0.5, 2.65, 6.0, 1.25, "Smart Guard AI Agent\nserver/smart-guard/  ·  enforce.ts fail-closed\nAllow / Step-up / Freeze"),
    (6.8, 2.65, 6.0, 1.25, "Financial engine\nExcel/CSV/PDF parse · P&L · leaks · Q&A on open file"),
    (0.5, 4.15, 3.9, 1.2, "PostgreSQL\nUsers · files · guard_decisions log"),
    (4.6, 4.15, 4.0, 1.2, "nokia-adapter.ts\nLive RapidAPI + structured [nac-live] audit logs"),
    (8.85, 4.15, 3.95, 1.2, "Nokia Network-as-Code\nCAMARA passthrough endpoints"),
]
for l, t, w, h, text in boxes:
    add_round(s, Inches(l), Inches(t), Inches(w), Inches(h), CARD)
    add_text(s, Inches(l + 0.18), Inches(t + 0.15), Inches(w - 0.35), Inches(h - 0.25), text, 12, False, WHITE)
add_text(s, Inches(0.55), Inches(5.55), Inches(12.2), Inches(1.2),
         "Security: NAC_API_KEY server-only · CSRF same-origin · scrypt passwords · guard logs with IP/UA · admin Smart Guard status dashboard.\n"
         "Repository: https://github.com/BITSANDBYTESDUDE/Smart-Profits",
         12, False, CYAN)
footer(s)

# ---------------------------------------------------------------------------
# 10 Live implementation proof
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "8. Live prototype evidence", "What judges can verify in code and logs")
card(s, Inches(0.5), Inches(1.45), Inches(6.05), Inches(2.35), "Live Nokia RapidAPI",
     "[nac-live] nacMode=live provider=Nokia transport=RapidAPI api=sim-swap status=200\n"
     "[nac-live] api=device-swap status=200\n"
     "[nac-live] api=number-verification status=401 → step-up (not fake success)\n"
     "Provider failures ≠ business rejection — distinct error messages in enforce.ts")
card(s, Inches(6.75), Inches(1.45), Inches(6.05), Inches(2.35), "Automated tests",
     "tests/nac-provider.test.ts — live mapper, auth/404 handling\n"
     "tests/smart-guard-policy.test.ts — Allow/Step-up/Freeze policy\n"
     "tests/backend-guard.test.ts — server enforcement fail-closed")
card(s, Inches(0.5), Inches(4.0), Inches(12.3), Inches(2.35), "Demo numbers (Nokia official sandbox MSISDNs)",
     "+99999991001 → clean path (register + analytics)\n"
     "+99999991000 → SIM/device swap → account FREEZE\n"
     "+99999991002 → location PARTIAL → step-up on file upload\n"
     "+99999991003 → location UNKNOWN (live API only)\n"
     "Step-up code: network challenge (6 digits) — shown in dev demo inbox; not spoofable SMS.")
footer(s)

# ---------------------------------------------------------------------------
# 11 Demo journey A/B/C
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "9. Demo walkthrough (3 minutes)", "Three journeys judges should see")
walk = [
    ("Journey A — Trusted merchant", "+99999991001 · new email · upload messy Excel · advisor: highest profit SKU · phantom vs real net profit · ALLOW"),
    ("Journey B — SIM swap attack", "+99999991000 · login or register · Smart Guard FREEZE · upload/export disabled · Arabic/EN explanation · SMS untrusted"),
    ("Journey C — Location mismatch", "+99999991002 · upload financial file · Location Verification PARTIAL · STEP-UP network code · file blocked until verified"),
]
for i, (title, body) in enumerate(walk):
    y = Inches(1.42) + i * Inches(1.22)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.08), CARD)
    add_text(s, Inches(0.75), y + Inches(0.1), Inches(12), Inches(0.32), title, 16, True, GOLD)
    add_text(s, Inches(0.75), y + Inches(0.45), Inches(11.8), Inches(0.55), body, 13, False, WHITE)
add_text(s, Inches(0.55), Inches(5.15), Inches(12.2), Inches(1.35),
         "Links for judges:\n"
         "GitHub: https://github.com/BITSANDBYTESDUDE/Smart-Profits\n"
         "Studio: https://bitsandbytesdude.vercel.app/\n"
         "Live demo URL: [deploy before submission — Railway/Vercel + NAC_API_KEY + PostgreSQL]",
         13, False, CYAN)
footer(s)

# ---------------------------------------------------------------------------
# 12 Screenshots
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "10. Product screenshots", "Working merchant SaaS — bilingual, dark/light, file-grounded AI")
shot_files = [
    ("01-login.png", "Login + Smart Guard"),
    ("02-dashboard.png", "Dashboard — health & P&L"),
    ("03-advisor.png", "Advisor — file-grounded Q&A"),
    ("04-data.png", "Files — messy Excel cleaned"),
    ("05-simulator.png", "What-if simulator"),
    ("06-guard-freeze.png", "Smart Guard freeze overlay"),
]
present = [(SHOTS / name, cap) for name, cap in shot_files if (SHOTS / name).exists()]
if not present:
    present = [(p, p.stem.replace("-", " ")) for p in sorted(SHOTS.glob("*.png"))[:6]]
if present:
    for i, (path, cap) in enumerate(present[:6]):
        col, row = i % 3, i // 3
        x = Inches(0.45) + col * Inches(4.25)
        y = Inches(1.38) + row * Inches(2.65)
        s.shapes.add_picture(str(path), x, y, Inches(4.05), Inches(2.15))
        add_text(s, x, y + Inches(2.18), Inches(4.05), Inches(0.3), cap, 10, True, CYAN)
else:
    add_text(s, Inches(0.55), Inches(2.5), Inches(12), Inches(2),
             "Run: node docs/hackathon/capture_screens.mjs (with npm run dev)\nto capture live screenshots into docs/hackathon/screenshots/",
             16, False, MUTED)
footer(s)

# ---------------------------------------------------------------------------
# 13 Differentiation
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "11. Differentiation", "Why Smart Profits is not another Excel chatbot")
headers = ["Typical Excel AI", "Smart Profits"]
rows = [
    ("Reads a spreadsheet", "Reads + diagnoses + bilingual Q&A on the open file"),
    ("Trusts SMS OTP", "SIM Swap + Number Verification via the network"),
    ("No location concept", "Location Verification binds uploads to the store"),
    ("Chatbot only", "AI Agent calls CAMARA tools and enforces freeze"),
    ("Generic SaaS", "Built for MENA merchants (Arabic RTL, messy AR/EN ledgers)"),
    ("Slide-only Open Gateway", "Live Nokia NaC integration with audit logs"),
]
table = s.shapes.add_table(7, 2, Inches(0.5), Inches(1.5), Inches(12.3), Inches(4.85)).table
table.columns[0].width = Inches(5.5)
table.columns[1].width = Inches(6.8)
for j, h in enumerate(headers):
    cell = table.cell(0, j)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(0x0B, 0x11, 0x20)
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            set_run(r, 12, True, CYAN)
for i, row in enumerate(rows, start=1):
    for j, val in enumerate(row):
        cell = table.cell(i, j)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD if i % 2 else RGBColor(0x14, 0x1C, 0x30)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_run(r, 12, j == 1, GOLD if j == 1 else WHITE)
footer(s)

# ---------------------------------------------------------------------------
# 14 Business model
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "12. Business model", "SaaS revenue + Open Gateway COGS + operator B2B2X")
plans = [
    ("Free trial · 7 days", "Full product — upload, diagnosis, advisor, Smart Guard demo. Converts to Pro."),
    ("Pro · $49 / month", "Single store — P&L, leaks, Q&A, telco-secured login/upload/export."),
    ("Business · $99 / month", "Multi-branch — tighter geofences, team seats, priority guard review."),
]
for i, (title, body) in enumerate(plans):
    card(s, Inches(0.5) + i * Inches(4.2), Inches(1.45), Inches(3.95), Inches(2.5), title, body)
add_round(s, Inches(0.5), Inches(4.2), Inches(12.3), Inches(2.35), CARD)
add_text(s, Inches(0.75), Inches(4.35), Inches(12), Inches(0.35), "Monetization logic for GSMA / Nokia / operators", 15, True, GOLD)
add_paras(s, Inches(0.75), Inches(4.75), Inches(11.8), Inches(1.65), [
    "Merchant subscription = primary revenue (SME FinTech SaaS).",
    "CAMARA API calls = security COGS — every login, upload and export consumes operator value.",
    "Future B2B2X: white-label Smart Guard for banks, wallets and marketplaces via Open Gateway.",
    "Incubated at Bitsandbytesdude — shipping product line, not a one-hack demo.",
], 13, WHITE, bullet=True)
footer(s)

# ---------------------------------------------------------------------------
# 15 Impact
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "13. Impact & why GSMA should care", "Real API consumer · Real regional problem · Real agent")
impacts = [
    ("For operators", "B2B FinTech that consumes Open Gateway APIs at scale — login, upload, export — not slide-only integration."),
    ("For merchants", "CFO-grade insight AND bank-like protection of the files that ARE the business."),
    ("For the region", "High WhatsApp/Excel commerce, high SIM-swap fraud, low formal accounting — exact Theme 4 conditions."),
    ("For 5G/4G", "Number Verification works because the phone is on the operator network — not because we sent another SMS."),
]
for i, (h, b) in enumerate(impacts):
    col, row = i % 2, i // 2
    card(s, Inches(0.5) + col * Inches(6.4), Inches(1.45) + row * Inches(2.45), Inches(6.05), Inches(2.2), h, b)
add_text(s, Inches(0.55), Inches(6.15), Inches(12.2), Inches(0.55),
         "Target outcome: Top 3 → showcase at MWC Doha 2026 (8–10 Nov) · prizes up to $10,000",
         13, True, GOLD)
footer(s)

# ---------------------------------------------------------------------------
# 16 Roadmap
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "14. Roadmap", "From hackathon prototype to MENA commercial product")
phases = [
    ("Now — Hackathon", "Live NaC · demo video · deploy public URL · pitch + GitHub README for judges."),
    ("Q4 2026", "NV OAuth on live operators · operator pilot · admin guard analytics · Pro billing."),
    ("2027", "Multi-branch geofences · B2B2X with MENA telco · marketplace integrations."),
    ("MWC Doha", "Live freeze demo on stage · executive pitch · operator meetings."),
]
for i, (h, b) in enumerate(phases):
    y = Inches(1.42) + i * Inches(1.18)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.02), CARD)
    add_text(s, Inches(0.75), y + Inches(0.1), Inches(3.2), Inches(0.32), h, 14, True, CYAN)
    add_text(s, Inches(3.5), y + Inches(0.12), Inches(9.2), Inches(0.75), b, 13, False, WHITE)
footer(s)

# ---------------------------------------------------------------------------
# 17 Team
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
chrome(s, "15. Team", "Gaza product inside a shipping software studio")
card(s, Inches(0.5), Inches(1.42), Inches(6.05), Inches(3.9), "Israa Nael Hamad — Founder / Product Lead",
     "University of Palestine, Gaza · Software Engineering (final year).\n\n"
     "Product vision, merchant UX, bilingual FinTech experience, Theme 4 narrative: phantom Excel profit vs SIM-swap takeover of the same files.\n\n"
     "Email: tsraathmd@gmail.com")
card(s, Inches(6.75), Inches(1.42), Inches(6.05), Inches(3.9), "Mir Kashif — Principal / Technical Co-Founder",
     "Bitsandbytesdude Software Agency.\n\n"
     "Full-stack platforms, AI agent workflows, cloud delivery. Next.js architecture, Nokia NaC integration, Smart Guard enforcement.\n\n"
     "Studio: bitsandbytesdude@gmail.com · https://bitsandbytesdude.vercel.app/")
add_round(s, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.15), CARD)
add_text(s, Inches(0.75), Inches(5.65), Inches(12), Inches(0.85),
         "Bitsandbytesdude — high-performance software, AI automation and product systems.\nRemote-first · MENA-rooted · Smart Profits is a core product line, not a side project.",
         13, False, WHITE)
footer(s)

# ---------------------------------------------------------------------------
# 18 Close / Ask
# ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, 0, 0, W, Inches(0.08), CYAN)
add_rect(s, 0, Inches(0.08), W, Inches(0.02), GOLD)
add_text(s, Inches(0.55), Inches(1.55), Inches(12), Inches(0.35), "THE ASK", 14, True, CYAN)
add_text(s, Inches(0.55), Inches(2.0), Inches(12.2), Inches(1.1), "Advance Smart Profits under Theme 4.", 38, True, WHITE)
add_paras(s, Inches(0.55), Inches(3.25), Inches(12), Inches(1.5), [
    "A real merchant SaaS — not slideware.",
    "A real AI Agent — Allow / Step-up / Freeze on every sensitive action.",
    "Real CAMARA APIs through Nokia Network-as-Code — with live logs judges can verify.",
    "Built for MENA — Arabic RTL, English, messy ledgers, SIM-swap-aware FinTech.",
], 16, WHITE, bullet=True)
badge_row(s, Inches(5.0), [
    "✓ CAMARA via Nokia NaC",
    "✓ AI Agent orchestration",
    "✓ Theme 4 FinTech",
    "✓ Live prototype",
])
add_text(s, Inches(0.55), Inches(5.75), Inches(12), Inches(0.85),
         "Thank you  ·  Questions?\n"
         "Israa Nael Hamad  ·  Mir Kashif  ·  Bitsandbytesdude  ·  University of Palestine, Gaza",
         15, False, GOLD)
footer(s)

out = ROOT / "Smart-Profits-HACKATHON-PITCH-DECK.pptx"
fallback = ROOT / "Smart-Profits-PITCH-DECK.pptx"
try:
    prs.save(out)
    print("saved", out)
except PermissionError:
    prs.save(fallback)
    print("locked; saved", fallback)
print("slides", page, "shots", len(present) if present else 0)
